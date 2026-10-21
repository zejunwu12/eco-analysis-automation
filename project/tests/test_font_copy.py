"""Tests for source-font copying at write time and font unification."""

import unittest

from openpyxl import Workbook
from openpyxl.styles import Font

from engine.font import apply_uniform_font
from engine.writer import copy_source_font


class CopySourceFontTests(unittest.TestCase):
    def setUp(self):
        self.source_book = Workbook()
        self.source_sheet = self.source_book.active
        self.target_book = Workbook()
        self.target_sheet = self.target_book.active

    def test_full_font_is_copied_like_manual_paste(self):
        self.source_sheet["A1"].font = Font(
            name="Calibri", size=14, bold=True, color="FFFF0000"
        )
        self.target_sheet["A1"].font = Font(name="宋体", size=11)

        copied = copy_source_font(
            self.source_sheet["A1"], self.target_sheet["A1"]
        )

        self.assertTrue(copied)
        target_font = self.target_sheet["A1"].font
        self.assertEqual(target_font.name, "Calibri")
        self.assertEqual(target_font.size, 14)
        self.assertTrue(target_font.bold)
        self.assertEqual(target_font.color.rgb, "FFFF0000")

    def test_default_style_source_leaves_target_font_untouched(self):
        self.target_sheet["A1"].font = Font(name="宋体", size=11)

        copied = copy_source_font(
            self.source_sheet["A1"], self.target_sheet["A1"]
        )

        self.assertFalse(copied)
        target_font = self.target_sheet["A1"].font
        self.assertEqual(target_font.name, "宋体")
        self.assertEqual(target_font.size, 11)
        self.assertIsNone(target_font.color)

    def test_copied_font_is_independent_of_source_workbook(self):
        self.source_sheet["A1"].font = Font(
            name="Calibri", size=14, color="FFFF0000"
        )

        copy_source_font(self.source_sheet["A1"], self.target_sheet["A1"])
        # 之后修改来源字体，不应影响已复制的目标单元格
        self.source_sheet["A1"].font = Font(name="Arial", size=9)

        target_font = self.target_sheet["A1"].font
        self.assertEqual(target_font.name, "Calibri")
        self.assertEqual(target_font.size, 14)
        self.assertEqual(target_font.color.rgb, "FFFF0000")


class ApplyUniformFontColorPreservationTests(unittest.TestCase):
    def _make_config(self):
        return {
            "font": {
                "enabled": True,
                "name": "宋体",
                "size": 11,
                "bold": False,
            },
            "reports": {
                "report7": {
                    "sheet_name": "报表7",
                    "data_start_col": "A",
                    "data_end_col": "C",
                    "sub_tables": [
                        {
                            "name": "子表",
                            "data_start_row": 4,
                            "data_end_row": 5,
                            "total_row": 6,
                        }
                    ],
                }
            },
        }

    def test_uniform_font_keeps_red_cells_red_and_normalizes_name_size(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "报表7"
        sheet["A4"] = "园区A"
        sheet["B4"] = "运营方A"
        sheet["C4"] = 100
        # 红色标注：最新调整数据
        for coordinate in ("A4", "B4", "C4"):
            sheet[coordinate].font = Font(
                name="Calibri", size=14, bold=True, color="FFFF0000"
            )
        # 普通黑色单元格
        sheet["A5"] = "园区B"
        sheet["A5"].font = Font(name="Calibri", size=14)

        modified = apply_uniform_font(workbook, self._make_config())

        self.assertGreater(modified, 0)
        for coordinate in ("A4", "B4", "C4"):
            font = sheet[coordinate].font
            self.assertEqual(font.name, "宋体")
            self.assertEqual(font.size, 11)
            self.assertFalse(font.bold)
            self.assertEqual(font.color.rgb, "FFFF0000")
        # 无颜色单元格统一后仍为默认（黑色）
        self.assertEqual(sheet["A5"].font.name, "宋体")
        self.assertEqual(sheet["A5"].font.size, 11)
        self.assertIsNone(sheet["A5"].font.color)

    def test_uniform_font_applies_report_override_and_keeps_color(self):
        config = self._make_config()
        config["font"]["report_overrides"] = {7: {"name": "微软雅黑", "size": 12}}
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "报表7"
        sheet["B4"].font = Font(name="Calibri", size=14, color="FFFF0000")

        apply_uniform_font(workbook, config)

        font = sheet["B4"].font
        self.assertEqual(font.name, "微软雅黑")
        self.assertEqual(font.size, 12)
        self.assertEqual(font.color.rgb, "FFFF0000")


class FixedReportLabelColumnFontTests(unittest.TestCase):
    """固定报表的标签列（集团名称/单位名称）也应统一字体。"""

    def _make_config(self):
        return {
            "font": {
                "enabled": True,
                "name": "宋体",
                "size": 11,
                "bold": False,
            },
            "reports": {
                "report4": {
                    "sheet_name": "报表4 各企业出租率",
                    "data_start_row": 4,
                    "data_end_row": 14,
                    "total_row": 15,
                    "a_col": "A",
                    "b_col": "B",
                    "data_start_col": "C",
                    "data_end_col": "K",
                }
            },
        }

    def _make_sheet(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "报表4 各企业出租率"
        # 标题与表头不属于统一范围
        sheet["A2"] = "文旅集团2026年第三季度各企业出租率"
        sheet["A2"].font = Font(name="方正小标宋简体", size=14, bold=True)
        for coordinate in ("A3", "B3", "C3"):
            sheet[coordinate] = "表头"
            sheet[coordinate].font = Font(name="宋体", size=11, bold=True)
        # 标签列中混杂的遗留字体（来自下发空表）
        sheet["A6"].font = Font(name="Calibri", size=11)
        sheet["A7"].font = Font(name="宋体", size=11)
        sheet["A8"] = "商业集团"
        sheet["A8"].font = Font(name="仿宋", size=11)
        sheet["B8"] = "商业集团"
        sheet["B8"].font = Font(name="仿宋", size=11)
        sheet["A14"] = "古城集团"
        sheet["A14"].font = Font(name="国标仿宋", size=11)
        sheet["B15"].font = Font(name="Calibri", size=11)
        sheet["A15"] = "合计"
        # 数据列：红色标注的最新调整数据
        sheet["C8"] = 127946.19
        sheet["C8"].font = Font(
            name="微软雅黑", size=14, bold=True, color="FFFF0000"
        )
        return workbook, sheet

    def test_label_columns_are_unified_with_data_area(self):
        workbook, sheet = self._make_sheet()

        modified = apply_uniform_font(workbook, self._make_config())

        # A4:K15 共 12 行 × 11 列
        self.assertEqual(modified, 12 * 11)
        for coordinate in ("A6", "A7", "A8", "B8", "A14", "B15", "A15"):
            font = sheet[coordinate].font
            self.assertEqual(font.name, "宋体", coordinate)
            self.assertEqual(font.size, 11, coordinate)
            self.assertFalse(font.bold, coordinate)
        # 数据列同样统一，且保留红色标注
        data_font = sheet["C8"].font
        self.assertEqual(data_font.name, "宋体")
        self.assertEqual(data_font.size, 11)
        self.assertFalse(data_font.bold)
        self.assertEqual(data_font.color.rgb, "FFFF0000")

    def test_title_and_header_rows_are_left_untouched(self):
        workbook, sheet = self._make_sheet()

        apply_uniform_font(workbook, self._make_config())

        title_font = sheet["A2"].font
        self.assertEqual(title_font.name, "方正小标宋简体")
        self.assertEqual(title_font.size, 14)
        self.assertTrue(title_font.bold)
        for coordinate in ("A3", "B3", "C3"):
            font = sheet[coordinate].font
            self.assertEqual(font.name, "宋体", coordinate)
            self.assertEqual(font.size, 11, coordinate)
            self.assertTrue(font.bold, coordinate)

    def test_label_column_falls_back_to_column_a_without_a_col_config(self):
        config = self._make_config()
        del config["reports"]["report4"]["a_col"]
        workbook, sheet = self._make_sheet()

        apply_uniform_font(workbook, config)

        self.assertEqual(sheet["A8"].font.name, "宋体")


if __name__ == "__main__":
    unittest.main()
